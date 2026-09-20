from docx import Document
from docx.shared import Pt

doc = Document()

title = doc.add_heading("Python Dictionaries — Practice Questions", level=1)
doc.add_heading("Beginner Level — Real Life Examples", level=2)

questions = [
    ("Contact Book",
     "Make dict `contacts` with 3 names as keys, phone numbers as values. Print one contact by name."),
    ("Add Entry",
     "Add new contact to `contacts` dict. Print updated dict."),
    ("Update Value",
     "Priya changed her phone number. Update existing key's value (don't add new key)."),
    ("Remove Entry",
     "Remove one contact using `.pop()`. Print removed value too."),
    ("Safe Remove",
     "Try remove contact not in dict using `.pop()` with default value (no crash)."),
    ("Check Existence",
     "Ask user for name, check `in` contacts dict before printing number. Print \"not found\" if missing."),
    ("Loop All",
     "Print all contacts as \"Name: Number\" using `.items()`."),
    ("Just Names",
     "Print only names using `.keys()`."),
    ("Just Numbers",
     "Print only numbers using `.values()`."),
    ("Student Grades",
     "Dict `grades` = subject : marks. Find subject with highest mark."),
    ("Grocery Inventory",
     "Dict `stock` = item : quantity. User buys item, reduce quantity by 1 using `.get()` for safe lookup."),
    ("Merge Dicts",
     "Two dicts `morning_shift`, `evening_shift` (employee : hours). Merge into one using `.update()`."),
    ("Nested Dict",
     'Student record: {"Amit": {"math": 90, "sci": 85}}. Access Amit\'s math marks.'),
    ("Count Words",
     "Given a sentence, count each word's frequency into a dict using `.get(word, 0) + 1` pattern."),
    ("Menu Price Lookup",
     'Restaurant menu dict item : price. Ask user for item name, use `.get()` with default "item not available".'),
]

for i, (title_text, body) in enumerate(questions, start=1):
    p = doc.add_paragraph()
    run = p.add_run(f"{i}. {title_text}")
    run.bold = True
    run.font.size = Pt(12)
    doc.add_paragraph(body)

doc.add_paragraph()
note = doc.add_paragraph()
note_run = note.add_run(
    "Covers: creating dicts, add/update, .pop(), in check, .items(), .keys(), "
    ".values(), .get(), .update(), nested dicts."
)
note_run.italic = True

doc.save("dictionary_practice_questions.docx")
print("saved")
